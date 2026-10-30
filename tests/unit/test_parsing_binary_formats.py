from __future__ import annotations

import io

import docx
import pytest
from packages.domain.errors import ParsingFailed
from rag.parsing.docx_parser import DocxParser
from rag.parsing.pdf_parser import PdfParser
from rag.parsing.registry import get_parser_for_mime_type
from reportlab.pdfgen import canvas


def _make_pdf(pages_text: list[str]) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    for text in pages_text:
        c.drawString(72, 720, text)
        c.showPage()
    c.save()
    return buffer.getvalue()


def _make_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    document = docx.Document()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
def test_pdf_parser_extracts_text_from_real_pdf() -> None:
    pdf_bytes = _make_pdf(["Hello from page one."])
    result = PdfParser().parse(pdf_bytes)
    assert "Hello from page one." in result.text
    assert result.metadata["page_count"] == 1


@pytest.mark.unit
def test_pdf_parser_tracks_page_boundaries_across_multiple_pages() -> None:
    pdf_bytes = _make_pdf(["Page one content.", "Page two content."])
    result = PdfParser().parse(pdf_bytes)

    assert result.metadata["page_count"] == 2
    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert result.pages[1].page_number == 2

    page_one_text = result.text[result.pages[0].character_start : result.pages[0].character_end]
    page_two_text = result.text[result.pages[1].character_start : result.pages[1].character_end]
    assert "Page one" in page_one_text
    assert "Page two" in page_two_text
    assert "Page two" not in page_one_text


@pytest.mark.unit
def test_pdf_parser_rejects_corrupted_content() -> None:
    with pytest.raises(ParsingFailed):
        PdfParser().parse(b"this is not a real pdf file at all")


@pytest.mark.unit
def test_docx_parser_extracts_paragraph_text() -> None:
    docx_bytes = _make_docx([("This is a normal paragraph.", None)])
    result = DocxParser().parse(docx_bytes)
    assert "This is a normal paragraph." in result.text


@pytest.mark.unit
def test_docx_parser_extracts_headings_with_correct_levels() -> None:
    docx_bytes = _make_docx(
        [
            ("Document Title", "Heading 1"),
            ("Some intro text.", None),
            ("A Section", "Heading 2"),
            ("Section body.", None),
        ]
    )
    result = DocxParser().parse(docx_bytes)
    assert [h.text for h in result.headings] == ["Document Title", "A Section"]
    assert [h.level for h in result.headings] == [1, 2]


@pytest.mark.unit
def test_docx_parser_rejects_corrupted_content() -> None:
    with pytest.raises(ParsingFailed):
        DocxParser().parse(b"not a real docx file")


@pytest.mark.unit
def test_registry_returns_correct_parser_per_mime_type() -> None:
    assert isinstance(get_parser_for_mime_type("application/pdf"), PdfParser)
    assert isinstance(
        get_parser_for_mime_type(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        DocxParser,
    )


@pytest.mark.unit
def test_registry_rejects_unregistered_mime_type() -> None:
    from packages.domain.errors import UnsupportedMimeType

    with pytest.raises(UnsupportedMimeType):
        get_parser_for_mime_type("application/x-executable")


@pytest.mark.unit
def test_end_to_end_registry_lookup_and_parse_for_pdf() -> None:
    pdf_bytes = _make_pdf(["Round trip via the registry."])
    parser = get_parser_for_mime_type("application/pdf")
    result = parser.parse(pdf_bytes)
    assert "Round trip via the registry." in result.text
